"""Génération d'un PPT explicatif éditable à partir des résultats du pipeline.

Jumeau de excel_report.build_excel : build_deck(results, output_path, site_name) produit un
.pptx PowerPoint natif (zones de texte, tableaux et formes éditables), que le consultant
récupère et intègre dans son propre deck client.

Style : charte WPP (navy), from scratch. Un template officiel (.potx) pourra être branché
en V2 sans changer l'API. Dépendance : python-pptx.
"""
from __future__ import annotations
import datetime as _dt

import pandas as pd
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# --- Charte WPP ---
NAVY = RGBColor(0x19, 0x1D, 0x63)
BLUE = RGBColor(0x3A, 0x3F, 0x8F)
ICE = RGBColor(0xCA, 0xDC, 0xFC)
LIGHT = RGBColor(0xEE, 0xF1, 0xFB)
GREY = RGBColor(0x6B, 0x6B, 0x6B)
DGREY = RGBColor(0x3C, 0x3C, 0x3C)
RED = RGBColor(0xC0, 0x39, 0x2A)
AMBER = RGBColor(0xC8, 0x90, 0x1F)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Calibri"
SERIF = "Cambria"


def _fmt(n):
    try:
        return f"{int(round(float(n))):,}".replace(",", " ")
    except Exception:
        return str(n)


def _text(slide, l, t, w, h, runs, size=14, bold=False, color=DGREY, align=PP_ALIGN.LEFT,
          font=FONT, italic=False, anchor=None, line_spacing=1.0):
    tb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    if anchor is not None:
        tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    if line_spacing:
        p.line_spacing = line_spacing
    if isinstance(runs, str):
        runs = [(runs, {})]
    for txt, o in runs:
        r = p.add_run()
        r.text = txt
        r.font.name = o.get("font", font)
        r.font.size = Pt(o.get("size", size))
        r.font.bold = o.get("bold", bold)
        r.font.italic = o.get("italic", italic)
        r.font.color.rgb = o.get("color", color)
    return tb


def _rect(slide, l, t, w, h, fill, line=None, rounded=True):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        Inches(l), Inches(t), Inches(w), Inches(h))
    shp.shadow.inherit = False
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(1)
    shp.text_frame.paragraphs[0].text = ""
    return shp


def _kicker_title(slide, kicker, title):
    _text(slide, 0.6, 0.42, 11, 0.3, kicker.upper(), size=12, bold=True, color=BLUE)
    _text(slide, 0.6, 0.74, 12.1, 0.8, title, size=27, bold=True, color=NAVY, font=SERIF)


def _foot(slide, site):
    _text(slide, 0.5, 7.06, 10, 0.3, f"WPP Media  |  Search  |  Audit de maillage interne  {site}",
          size=9, color=GREY)


# --------------------------------------------------------------------------
def _metrics(results):
    p = results["pages"]
    rep = results["report"]
    links = results.get("links")
    m = {}
    n = max(int(rep.get("pages_total", len(p))), 1)
    m["pages"] = rep.get("pages_total", len(p))
    m["liens_valides"] = rep.get("liens_valides", "")
    m["orph"] = rep.get("pages_orphelines", int(p.get("orpheline", pd.Series(dtype=bool)).sum()) if "orpheline" in p else 0)
    m["orph_pct"] = m["orph"] / n * 100
    if links is not None and "position" in links.columns:
        pos = links["position"].value_counts()
        tot = max(int(pos.sum()), 1)
        m["boiler_pct"] = (pos.get("Footer", 0) + pos.get("Header", 0)) / tot * 100
        m["content_pct"] = pos.get("Content", 0) / tot * 100
    else:
        m["boiler_pct"] = m["content_pct"] = None
    quad = p["quadrant_code"].value_counts().to_dict() if "quadrant_code" in p else {}
    m["quad"] = {q: int(quad.get(q, 0)) for q in ["Q1", "Q2", "Q3", "Q4"]}
    m["js_share"] = results.get("_js_share")  # hook V2 (crawl JS)
    return m


def build_deck(results: dict, output_path: str, site_name: str = "Client") -> str:
    p = results["pages"]
    m = _metrics(results)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # ---- S1 cover ----
    s = prs.slides.add_slide(blank)
    _rect(s, -0.1, -0.1, 13.6, 7.7, NAVY, rounded=False)
    _text(s, 0.7, 0.8, 11, 0.4, "WPP MEDIA  |  SEARCH", size=13, bold=True, color=ICE)
    _text(s, 0.7, 2.5, 12, 0.9, "Audit de maillage interne", size=42, bold=True, color=WHITE, font=SERIF)
    _text(s, 0.7, 3.45, 12, 0.9, site_name, size=42, bold=True, color=ICE, font=SERIF)
    vol = f"{_fmt(m['pages'])} pages"
    if m["liens_valides"] != "":
        vol += f"  ·  {_fmt(m['liens_valides'])} liens internes valides"
    _text(s, 0.72, 4.7, 12, 0.5, f"{vol}  ·  {_dt.date.today().strftime('%d/%m/%Y')}",
          size=15, color=ICE)

    # ---- S2 en bref (callouts) ----
    s = prs.slides.add_slide(blank)
    _kicker_title(s, "1. En bref", "L'essentiel en quatre chiffres")
    cards = []
    if m["boiler_pct"] is not None:
        cards.append((f"{m['boiler_pct']:.0f} %", "du maillage dans le menu et le pied de page", NAVY))
        cards.append((f"{m['content_pct']:.0f} %", "seulement de liens dans le contenu", RED if m["content_pct"] < 25 else BLUE))
    cards.append((_fmt(m["orph"]), f"pages orphelines ({m['orph_pct']:.0f} %)", AMBER))
    cards.append((str(m["quad"]["Q1"]), "pages prioritaires (Q1) à renforcer", BLUE))
    cards = cards[:4]
    cw, gap, x0 = 2.86, 0.3, 0.6
    for i, (num, lab, col) in enumerate(cards):
        x = x0 + i * (cw + gap)
        _rect(s, x, 2.5, cw, 2.4, LIGHT, line=ICE)
        _text(s, x, 2.85, cw, 0.95, num, size=36, bold=True, color=col, align=PP_ALIGN.CENTER, font=SERIF)
        _text(s, x + 0.2, 3.85, cw - 0.4, 0.9, lab, size=12, color=DGREY, align=PP_ALIGN.CENTER, line_spacing=1.02)
    _text(s, 0.6, 5.4, 12.1, 0.6, [(
        "Le levier n°1 : déplacer de l'énergie du menu/pied de page vers du lien éditorial "
        "qui pointe vraiment les pages à enjeu.", {"italic": True, "color": NAVY, "size": 14})])
    _foot(s, site_name)

    # ---- S3 matrice 2x2 ----
    s = prs.slides.add_slide(blank)
    _kicker_title(s, "2. Où agir", "Le diagnostic : deux axes, quatre quadrants")
    _text(s, 1.15, 1.62, 5, 0.3, "↓  DÉFICIT DE MAILLAGE", size=11, bold=True, color=GREY)
    _text(s, 7.0, 1.62, 5.8, 0.3, "POTENTIEL BUSINESS  →", size=11, bold=True, color=GREY, align=PP_ALIGN.RIGHT)
    gx, gy, qw, qh, qg = 1.15, 1.98, 5.72, 2.18, 0.18
    quads = [
        (0, 0, "Q4", "Ne rien faire", m["quad"]["Q4"], RGBColor(0xF2, 0xF2, 0xF2), DGREY, "Faible potentiel · fort déficit"),
        (0, 1, "Q1", "Prioritaire", m["quad"]["Q1"], RGBColor(0xF7, 0xD9, 0xD5), RED, "Fort potentiel · fort déficit"),
        (1, 0, "Q3", "Sur-maillage", m["quad"]["Q3"], RGBColor(0xFC, 0xEF, 0xCF), AMBER, "Faible potentiel · bon maillage"),
        (1, 1, "Q2", "Hors scope maillage", m["quad"]["Q2"], LIGHT, BLUE, "Fort potentiel · bon maillage"),
    ]
    for r, c, code, name, val, fill, tc, sub in quads:
        x, y = gx + c * (qw + qg), gy + r * (qh + qg)
        _rect(s, x, y, qw, qh, fill, line=RGBColor(0xD8, 0xDB, 0xE8))
        _text(s, x + 0.3, y + 0.22, 2, 0.5, code, size=22, bold=True, color=tc, font=SERIF)
        _text(s, x + 0.3, y + 0.78, qw - 2.3, 0.5, name, size=15, bold=True, color=DGREY)
        _text(s, x + 0.3, y + 1.26, qw - 2.2, 0.6, sub, size=11, color=GREY)
        _text(s, x + qw - 2.15, y + 0.5, 1.9, 1.1, str(val), size=38, bold=True, color=tc,
              align=PP_ALIGN.RIGHT, font=SERIF)
    _text(s, 1.15, 6.5, 11.5, 0.4, [("Lecture : ", {"bold": True, "color": NAVY}),
          ("renforcer les Q1 (prescrire des liens), réallouer le maillage des Q3 vers ces Q1.",
           {"color": DGREY})], size=13, italic=True)
    _foot(s, site_name)

    # ---- S4 quick wins (top pages problématiques par importance) ----
    s = prs.slides.add_slide(blank)
    _kicker_title(s, "3. Priorités", "Pages à fort trafic, maillage défaillant")
    rows_data = _quick_wins(results)
    if rows_data:
        _text(s, 0.6, 1.5, 12.1, 0.4, "Pages qui pèsent le plus (clics) et dont le maillage ou les ancres posent problème.",
              size=14, color=DGREY)
        _qw_table(s, rows_data)
    else:
        _text(s, 0.6, 1.7, 12.1, 0.5, "Ajouter les données Search Console pour prioriser les pages par trafic réel.",
              size=14, italic=True, color=GREY)
    _foot(s, site_name)

    # ---- S5 flag JS (optionnel, si crawl JS) ----
    if m["js_share"] is not None:
        s = prs.slides.add_slide(blank)
        _kicker_title(s, "4. Vigilance", "Une partie du maillage dépend du JavaScript")
        _rect(s, 0.6, 2.6, 12.13, 1.6, RGBColor(0xFC, 0xEF, 0xCF), line=RGBColor(0xEA, 0xD9, 0xA8))
        _text(s, 0.95, 2.85, 11.4, 0.9, [(f"{m['js_share']:.0f} % ", {"bold": True, "color": AMBER, "size": 28}),
              ("des liens internes contextuels n'existent qu'après exécution du JavaScript.",
               {"color": DGREY, "size": 17})], anchor=MSO_ANCHOR.MIDDLE)
        _text(s, 0.6, 4.5, 12.1, 0.6, "Enjeu : un moteur qui rend mal le JavaScript peut rater ces liens. On sait les repérer.",
              size=14, color=NAVY)
        _foot(s, site_name)

    # ---- S6 plan d'action ----
    s = prs.slides.add_slide(blank)
    _rect(s, -0.1, -0.1, 13.6, 7.7, NAVY, rounded=False)
    _text(s, 0.7, 0.9, 11, 0.4, "EN RÉSUMÉ", size=13, bold=True, color=ICE)
    _text(s, 0.7, 1.75, 12, 0.8, "Ce qu'on fait pour corriger", size=32, bold=True, color=WHITE, font=SERIF)
    steps = [
        ("1", "Corriger la navigation", "Recibler les liens du menu et du pied de page vers les URL finales (pas les redirections)."),
        ("2", "Densifier le maillage contextuel", "Poser les liens prescrits en contenu vers les pages prioritaires, avec les bonnes ancres."),
        ("3", "Traiter les pages orphelines", "Reconnecter les pages sans lien entrant, en commençant par celles à fort trafic."),
        ("4", "Réallouer le PageRank", "Alléger les liens vers les pages sans enjeu (Q3) au profit des pages prioritaires."),
    ]
    y = 2.9
    for num, ttl, desc in steps:
        c = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.9), Inches(y), Inches(0.55), Inches(0.55))
        c.shadow.inherit = False
        c.fill.solid(); c.fill.fore_color.rgb = ICE; c.line.fill.background()
        _tf = c.text_frame; _tf.word_wrap = False
        _pp = _tf.paragraphs[0]; _pp.alignment = PP_ALIGN.CENTER
        _rn = _pp.add_run(); _rn.text = num
        _rn.font.name = SERIF; _rn.font.size = Pt(22); _rn.font.bold = True; _rn.font.color.rgb = NAVY
        _text(s, 1.7, y - 0.05, 11, 0.45, ttl, size=18, bold=True, color=WHITE)
        _text(s, 1.7, y + 0.4, 10.9, 0.5, desc, size=13, color=ICE)
        y += 1.02

    prs.save(output_path)
    return output_path


def _quick_wins(results, top=5):
    """Top pages problématiques triées par clics GSC. Retourne liste de (page, clics, probleme, action)."""
    p = results["pages"]
    flags = results.get("flags")
    if flags is None or not len(flags) or "clics_gsc" not in p.columns:
        return []
    imp = p[["url", "clics_gsc"]]
    f = flags.merge(imp, on="url", how="left")
    f = f[pd.to_numeric(f["clics_gsc"], errors="coerce").fillna(0) > 0]
    if not len(f):
        return []
    # garder le problème le plus grave par page, puis trier par clics
    grav = {"Élevée": 0, "Moyenne": 1, "Faible": 2}
    f["_g"] = f["gravite"].map(grav).fillna(3)
    f = f.sort_values(["_g", "clics_gsc"], ascending=[True, False]).drop_duplicates("url")
    f = f.sort_values("clics_gsc", ascending=False).head(top)
    actions = {
        "Page orpheline": "Reconnecter depuis les pages fortes du même thème",
        "Ancre pauvre": "Réécrire les ancres sur les mots-clés cibles",
        "Ancre sur-homogène": "Diversifier les ancres entrantes",
        "Cannibalisation": "Clarifier la page cible par mot-clé",
        "Cluster isolé": "Ouvrir le cluster vers les catégories proches",
    }
    out = []
    for _, r in f.iterrows():
        url = str(r["url"]).replace("https://", "").replace("http://", "")
        # raccourcir le domaine
        if "/" in url:
            url = "/" + url.split("/", 1)[1]
        out.append((url[:46], _fmt(r["clics_gsc"]), str(r["type"]), actions.get(r["type"], "Corriger")))
    return out


def _qw_table(slide, rows_data):
    headers = ["Page", "Clics GSC", "Problème", "Action"]
    widths = [Inches(4.4), Inches(1.5), Inches(2.6), Inches(3.6)]
    nrows = len(rows_data) + 1
    tbl_shape = slide.shapes.add_table(nrows, 4, Inches(0.6), Inches(2.05),
                                       Inches(12.1), Inches(0.5 * nrows))
    table = tbl_shape.table
    for i, w in enumerate(widths):
        table.columns[i].width = w
    for j, htxt in enumerate(headers):
        cell = table.cell(0, j)
        cell.fill.solid(); cell.fill.fore_color.rgb = NAVY
        pr = cell.text_frame.paragraphs[0]; pr.alignment = PP_ALIGN.CENTER if j == 1 else PP_ALIGN.LEFT
        run = pr.add_run(); run.text = htxt
        run.font.name = FONT; run.font.size = Pt(12); run.font.bold = True; run.font.color.rgb = WHITE
    for i, row in enumerate(rows_data, start=1):
        for j, val in enumerate(row):
            cell = table.cell(i, j)
            cell.fill.solid(); cell.fill.fore_color.rgb = WHITE if i % 2 else LIGHT
            pr = cell.text_frame.paragraphs[0]; pr.alignment = PP_ALIGN.CENTER if j == 1 else PP_ALIGN.LEFT
            run = pr.add_run(); run.text = str(val)
            run.font.name = FONT; run.font.size = Pt(11)
            run.font.color.rgb = RED if j == 1 else DGREY
            run.font.bold = (j == 1)
